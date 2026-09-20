#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将 Excel(.xlsx) 转换为用于前端预览的 HTML 表格片段(每个 sheet 一个 table)。
用法: python3 xlsx_to_html.py <file.xlsx>
输出纯 HTML 片段(无 <html>/<body>),直接注入预览弹窗的 v-html 容器。
"""
import sys
import html
from openpyxl import load_workbook


def sheet_to_table(ws):
    rows = list(ws.iter_rows(values_only=True))
    parts = [f'<h3 style="margin:14px 0 6px;font-size:15px;">{html.escape(str(ws.title))}</h3>']
    parts.append(
        '<table border="1" cellspacing="0" cellpadding="6" '
        'style="border-collapse:collapse;width:100%;font-size:13px;'
        'background:#fff;">'
    )
    if not rows:
        parts.append('<tr><td style="color:#999;padding:8px;">（空表）</td></tr>')
    for r in rows:
        parts.append('<tr>')
        for c in r:
            v = '' if c is None else str(c)
            # 数字右对齐,其余默认左对齐
            align = 'text-align:right;' if isinstance(c, (int, float)) else ''
            parts.append(
                f'<td style="border:1px solid #e0e0e0;padding:4px 8px;{align}">'
                f'{html.escape(v)}</td>'
            )
        parts.append('</tr>')
    parts.append('</table>')
    return ''.join(parts)


def main():
    if len(sys.argv) < 2:
        print('<p style="color:#f56c6c;">未提供 Excel 文件路径</p>', file=sys.stderr)
        sys.exit(2)
    src = sys.argv[1]
    try:
        wb = load_workbook(src, data_only=True)
    except Exception as e:
        print(f'<p style="color:#f56c6c;">无法读取 Excel 文件: {html.escape(str(e))}</p>')
        sys.exit(0)

    out = []
    for ws in wb.worksheets:
        out.append(sheet_to_table(ws))
    wb.close()
    print('\n'.join(out))


if __name__ == '__main__':
    main()
