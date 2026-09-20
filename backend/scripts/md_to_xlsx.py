#!/usr/bin/env python3
# 将 Markdown 文档中的表格转换为 Excel(.xlsx)
# 用法: python3 md_to_xlsx.py <input.md> <output.xlsx>
# 每个连续表格块 -> 一个 sheet，sheet 名取该表格前最近的 #~#### 标题(截断 31 字符)
import sys, re
from openpyxl import Workbook
from openpyxl.styles import Font

# openpyxl 不允许出现在 sheet 标题中的字符(ASCII)
_FORBIDDEN = set('\\/:*?[]')
# 常见全角/变体斜杠(部分 Markdown 会用)，一并替换为 '-'
_VARIANT_SLASH = {'／', '⁄', '∕', '╱', '⧸', '⟋'}


def _sanitize_title(name):
    out = []
    for c in name:
        if c in _FORBIDDEN or c in _VARIANT_SLASH:
            out.append('-')
        else:
            out.append(c)
    s = ''.join(out).strip()
    s = s[:31]
    return s or 'Sheet'


def md_to_xlsx(md_path, xlsx_path):
    try:
        with open(md_path, encoding='utf-8') as f:
            md = f.read()
    except Exception as e:
        print(f"read error: {e}")
        sys.exit(1)

    lines = md.split('\n')
    wb = Workbook()
    wb.remove(wb.active)
    i = 0
    table_idx = 0
    last_heading = 'Sheet'
    used_titles = set()
    while i < len(lines):
        line = lines[i]
        h = re.match(r'^#{1,4}\s+(.+)$', line)
        if h:
            last_heading = h.group(1).strip()[:31]
        # 表格块:连续以 | 起始并以 | 结尾的非空行
        if line.strip().startswith('|') and line.strip().endswith('|') and '|' in line[1:-1]:
            block = [line]
            j = i + 1
            while j < len(lines) and lines[j].strip().startswith('|') and lines[j].strip().endswith('|'):
                block.append(lines[j])
                j += 1
            # 第二行须为分隔行(全 - : | 空格)
            if len(block) >= 2 and re.match(r'^[\|\s:\-]+$', block[1].strip()):
                table_idx += 1
                base = last_heading if last_heading != 'Sheet' else f'table_{table_idx}'
                title = _sanitize_title(base)
                # 保证 sheet 名唯一,避免 openpyxl 抛 "already exists"
                if title in used_titles:
                    k = 2
                    while f'{title}_{k}' in used_titles:
                        k += 1
                    title = f'{title}_{k}'
                used_titles.add(title)
                ws = wb.create_sheet(title=title)
                for r in block:
                    cells = [c.strip() for c in r.strip().strip('|').split('|')]
                    if re.match(r'^[\s:\-]+$', ''.join(cells)):
                        continue  # 跳过分隔行
                    ws.append(cells)
                # 表头加粗
                for c in ws[1]:
                    c.font = Font(bold=True)
                # 列宽自适应
                for col in ws.columns:
                    maxlen = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                    ws.column_dimensions[col[0].column_letter].width = min(max(maxlen + 2, 10), 60)
                i = j
                continue
        i += 1

    if wb.sheetnames:
        wb.save(xlsx_path)
        print("OK sheets:", ",".join(wb.sheetnames))
    else:
        # 兜底:无表格时仍生成一张"内容"表,保证 Excel 交付物必然存在(全量放开三格式需要)
        title = '内容'
        if title in used_titles:
            title = '内容_2'
        used_titles.add(title)
        ws = wb.create_sheet(title=title)
        for line in md.split('\n'):
            s = line.strip()
            if not s:
                continue
            if re.match(r'^[\|\s:\-]+$', s):  # 跳过分隔噪音行
                continue
            if s.startswith('|') and s.endswith('|') and '|' in s[1:-1]:
                s = s[1:-1].strip()
            ws.append([s])
        ws.column_dimensions['A'].width = 80
        wb.save(xlsx_path)
        print("OK sheets(兜底): 内容")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("usage: md_to_xlsx.py <input.md> <output.xlsx>")
        sys.exit(1)
    md_to_xlsx(sys.argv[1], sys.argv[2])
