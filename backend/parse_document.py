#!/usr/bin/env python3
"""
通用文档解析器 - 支持 docx/xlsx/pdf/md/txt/csv
用法: python3 parse_document.py <file_path>
输出: 解析后的纯文本内容到 stdout
退出码: 0=成功, 非0=失败
"""
import sys
import os

def parse_docx(file_path):
    """解析 .docx 文件"""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return '\n'.join(paragraphs)
    except Exception as e:
        raise RuntimeError(f"DOCX解析失败: {e}")

def parse_xlsx(file_path):
    """解析 .xlsx 文件"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets_content = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines = [f"[Sheet: {sheet_name}]"]
            for row in ws.iter_rows(values_only=True):
                # 过滤掉全为 None 的行
                row_vals = [str(v) if v is not None else '' for v in row]
                row_text = ' | '.join(v for v in row_vals if v)
                if row_text:
                    lines.append(row_text)
            if len(lines) > 1:  # 有内容才加
                sheets_content.append('\n'.join(lines))
        return '\n\n'.join(sheets_content)
    except Exception as e:
        raise RuntimeError(f"XLSX解析失败: {e}")

def parse_pdf(file_path):
    """解析 .pdf 文件"""
    try:
        import pdfplumber
        pages_text = []
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(f"[Page {i+1}]\n{text.strip()}")
        result = '\n\n'.join(pages_text)
        if not result.strip():
            raise RuntimeError("PDF内容为空或无法提取")
        return result
    except Exception as e:
        raise RuntimeError(f"PDF解析失败: {e}")

def parse_plain_text(file_path):
    """解析纯文本文件 (.txt/.md/.csv)"""
    encodings = ['utf-8', 'gbk', 'gb2312', 'gb18030', 'latin-1']
    for enc in encodings:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise RuntimeError(f"文本文件编码不支持: {file_path}")

def main():
    if len(sys.argv) < 2:
        print("用法: python3 parse_document.py <file_path>", file=sys.stderr)
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    if not os.path.exists(file_path):
        print(f"文件不存在: {file_path}", file=sys.stderr)
        sys.exit(1)
    
    ext = os.path.splitext(file_path)[1].lower()
    
    try:
        if ext == '.docx':
            content = parse_docx(file_path)
        elif ext == '.xlsx':
            content = parse_xlsx(file_path)
        elif ext == '.pdf':
            content = parse_pdf(file_path)
        elif ext in ('.txt', '.md', '.csv'):
            content = parse_plain_text(file_path)
        else:
            # 尝试当作纯文本处理
            content = parse_plain_text(file_path)
        
        if not content or not content.strip():
            print(f"文档内容为空: {file_path}", file=sys.stderr)
            sys.exit(1)
        
        print(content)
        sys.exit(0)
        
    except Exception as e:
        print(f"解析错误: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == '__main__':
    main()
