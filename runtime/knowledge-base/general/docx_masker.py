# _*_ coding: utf-8 _*_
# 对 docx 文档中的文本内容进行数据脱敏（正则规则 + 敏感词表，只对命中部分打码）。

import os
import re
import argparse
from docx import Document

# ========= 正则规则区域，与 image_masker 一致 =========

# 手机号（11 位，以 1 开头）
PHONE_PATTERN = re.compile(r'\b1\d{10}\b')

# 身份证号（15 或 18 位，简单版）
ID_PATTERN = re.compile(r'\b\d{15}(\d{2}[0-9Xx])?\b')

# 邮箱
EMAIL_PATTERN = re.compile(r'\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}\b')

# 银行卡 / 长数字串（一般 12 位以上）
LONG_NUMBER_PATTERN = re.compile(r'\b\d{12,}\b')

# 敏感词表路径：与脚本同目录下的 敏感词表.txt
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SENSITIVE_WORDS_FILE = os.path.join(_SCRIPT_DIR, "敏感词表.txt")


def mask_phone(match: re.Match) -> str:
    """手机号脱敏：138****1234"""
    s = match.group()
    return s[:3] + "****" + s[-4:]


def mask_id(match: re.Match) -> str:
    """身份证脱敏：前 4 + 中间全星 + 后 2"""
    s = match.group()
    if len(s) < 6:
        return "*" * len(s)
    return s[:4] + "*" * (len(s) - 6) + s[-2:]


def mask_email(match: re.Match) -> str:
    """邮箱脱敏：a***@example.com"""
    s = match.group()
    name, domain = s.split("@", 1)
    if len(name) <= 1:
        return "*" * len(name) + "@" + domain
    return name[0] + "***@" + domain


def mask_long_number(match: re.Match) -> str:
    """
    长数字串脱敏：保留前 4 + 中间全星 + 后 4
    例：6222021234567890 -> 6222********7890
    """
    s = match.group()
    if len(s) <= 8:
        return "*" * len(s)
    return s[:4] + "*" * (len(s) - 8) + s[-4:]


def load_sensitive_words(file_path: str = None) -> list:
    """
    从敏感词表文件加载敏感词列表。
    每行一个词，# 开头和空行忽略。文件不存在或读失败时返回空列表。
    """
    path = file_path or SENSITIVE_WORDS_FILE
    words = []
    if not os.path.isfile(path):
        return words
    try:
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#"):
                    words.append(line)
    except Exception:
        pass
    return words


def mask_text(text: str, sensitive_words: list = None) -> str:
    """
    对一段普通文本做脱敏：正则命中按规则替换；敏感词表命中词首尾保留、中间打星（如 中国移动 -> 中**动）。
    """
    if not text:
        return text

    # 先按正则做格式化脱敏（手机号、身份证、邮箱、长数字）
    text = EMAIL_PATTERN.sub(mask_email, text)
    text = ID_PATTERN.sub(mask_id, text)
    text = PHONE_PATTERN.sub(mask_phone, text)
    text = LONG_NUMBER_PATTERN.sub(mask_long_number, text)

    # 敏感词表：只对命中的词打码，首尾保留、中间打星，如 中国移动 -> 中**动
    if sensitive_words:
        for w in sensitive_words:
            if not w:
                continue
            if len(w) == 1:
                replacement = "*"
            elif len(w) == 2:
                replacement = w[0] + "*"
            else:
                replacement = w[0] + "*" * (len(w) - 2) + w[-1]
            text = text.replace(w, replacement)

    return text


def process_paragraph(paragraph, sensitive_words: list):
    """处理一个段落内的所有 runs 文本"""
    for run in paragraph.runs:
        original = run.text
        masked = mask_text(original, sensitive_words)
        if masked != original:
            run.text = masked


def process_table(table, sensitive_words: list):
    """递归处理表格中的所有单元格（包括嵌套表格）"""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                process_paragraph(paragraph, sensitive_words)
            for nested_table in cell.tables:
                process_table(nested_table, sensitive_words)


def mask_docx(docx_path: str) -> str:
    """
    对单个 docx 文档做脱敏，返回新文件路径。
    新文件命名为：原名_脱敏.docx
    使用 代码/敏感词表.txt，只对命中的敏感词和正则匹配处打码。
    """
    if not os.path.isfile(docx_path):
        raise FileNotFoundError(f"文件不存在：{docx_path}")

    doc = Document(docx_path)
    sensitive_words = load_sensitive_words()

    for paragraph in doc.paragraphs:
        process_paragraph(paragraph, sensitive_words)

    for table in doc.tables:
        process_table(table, sensitive_words)

    # 生成新文件名
    dir_name = os.path.dirname(docx_path)
    base_name = os.path.basename(docx_path)
    name_without_ext, ext = os.path.splitext(base_name)

    new_name = f"{name_without_ext}_脱敏{ext}"
    new_path = os.path.join(dir_name, new_name)

    doc.save(new_path)
    return new_path


def main():
    # parser = argparse.ArgumentParser(
    #     description="对 docx 文档中的文本内容进行数据脱敏（正则规则 + 敏感词表，只对命中部分打码）。"
    # )
    # parser.add_argument(
    #     "docx_path",
    #     help="要脱敏的 docx 文件路径",
    # )

    # args = parser.parse_args()
    # new_path = mask_docx(args.docx_path)
    # print("脱敏完成：")
    # print(f"- 原文件：{args.docx_path}")
    # print(f"- 新文件：{new_path}")
    docx_path = r"D:\CursorProjectManage\Caixun_Pinzhibu_Zhinengkefu_project\需求文档\V0.0.1迭代\v0.0.1.docx"
    new_path = mask_docx(docx_path)
    print("脱敏完成：")
    print(f"- 原文件：{docx_path}")
    print(f"- 新文件：{new_path}")



if __name__ == "__main__":
    main()